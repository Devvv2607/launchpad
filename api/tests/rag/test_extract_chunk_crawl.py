from __future__ import annotations

import io

import httpx
import pytest

from launchpad.rag import crawl as crawl_mod
from launchpad.rag.chunk import MAX_TOKENS, MIN_TOKENS, chunk_sections, est_tokens
from launchpad.rag.crawl import CrawlError, assert_public_url, crawl
from launchpad.rag.extract import Section, from_docx, from_html, from_pdf, from_text

PARA = (
    "Chai & Chapter is a bookshop cafe in Bandra. We brew cutting chai over charcoal, "
    "serve filter coffee in steel tumblers and host a reading circle every Sunday. "
)


def test_markdown_sections_keep_headings() -> None:
    md = "# Menu\nCutting chai\nFilter coffee\n\n## Events\nSunday reading circle\n"
    sections = from_text(md, "brand.md", markdown=True)
    assert [(s.heading, s.text) for s in sections] == [
        ("Menu", "Cutting chai\nFilter coffee"),
        ("Events", "Sunday reading circle"),
    ]


def test_html_strips_boilerplate() -> None:
    html = """<html><head><title>Chai & Chapter</title></head><body>
      <nav><a href="/menu">Menu</a></nav>
      <div class="cookie-banner">We use cookies</div>
      <main><h1>Our story</h1><p>Started in 2019 by two book lovers.</p>
        <h2>Menu</h2><ul><li>Cutting chai</li><li>Bun maska</li></ul></main>
      <footer>© 2026 All rights reserved</footer></body></html>"""
    sections, links, title = from_html(html, "https://chai.example/")
    text = " ".join(s.text for s in sections)
    assert title == "Chai & Chapter"
    assert "two book lovers" in text and "Bun maska" in text
    assert "cookies" not in text and "rights reserved" not in text
    assert [s.heading for s in sections] == ["Our story", "Menu"]
    assert "/menu" in links


def test_docx_headings_and_tables() -> None:
    import docx

    d = docx.Document()
    d.add_heading("Brand voice", level=1)
    d.add_paragraph("Warm, bookish and a little cheeky.")
    d.add_heading("Offers", level=2)
    d.add_paragraph("Monsoon special: 20% off cutting chai.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Chai", "Rs 40"
    buf = io.BytesIO()
    d.save(buf)
    sections = from_docx(buf.getvalue(), "brand.docx")
    assert [s.heading for s in sections] == ["Brand voice", "Offers", "Table"]
    assert "Rs 40" in sections[-1].text


def _pdf(pages: list[str]) -> bytes:
    """Hand-rolled minimal PDF with one text line per page (no extra dependencies)."""
    objs: list[bytes] = []
    kids = " ".join(f"{3 + i * 2} 0 R" for i in range(len(pages)))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_ref = 3 + len(pages) * 2
    for i, text in enumerate(pages):
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + i * 2} 0 R "
            f"/Resources << /Font << /F1 {font_ref} 0 R >> >> >>".encode()
        )
        objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return bytes(out)


def test_pdf_pages_are_cited() -> None:
    sections = from_pdf(_pdf(["Our story begins in Bandra", "Menu cutting chai"]), "brand.pdf")
    assert [(s.page, "Bandra" in s.text or "chai" in s.text) for s in sections] == [
        (1, True),
        (2, True),
    ]


def test_chunks_respect_size_band_overlap_and_headings() -> None:
    long_text = "\n".join(PARA * 3 for _ in range(12))
    sections = [
        Section(long_text, "About us", source="brand.md"),
        Section("Open 8am-11pm daily.", "Hours", source="brand.md"),
    ]
    chunks = chunk_sections(sections)
    about = [c for c in chunks if c.heading == "About us"]
    assert len(about) >= 3
    assert all(c.tokens <= MAX_TOKENS + 20 for c in chunks)
    assert all(c.tokens >= MIN_TOKENS // 2 for c in about[:-1])
    # Overlap: consecutive chunks share text.
    tail = about[0].text.splitlines()[-1]
    assert tail in about[1].text
    # Every chunk is prefixed with (and attributed to) exactly one heading.
    assert all(c.text.startswith(c.heading or "") for c in chunks)
    hours = [c for c in chunks if c.heading == "Hours"]
    assert len(hours) == 1 and "8am" in hours[0].text
    assert est_tokens("x" * 400) == 100


@pytest.fixture
def fake_dns(monkeypatch: pytest.MonkeyPatch) -> dict[str, set[str]]:
    table = {
        "chai.example": {"93.184.216.34"},
        "evil.example": {"10.0.0.5"},
        "localhost": {"127.0.0.1"},
    }

    async def resolve(host: str, port: int) -> set[str]:
        return table.get(host, {"93.184.216.34"})

    monkeypatch.setattr(crawl_mod, "_resolve", resolve)
    return table


@pytest.mark.parametrize(
    "url",
    ["http://localhost/admin", "http://evil.example/", "file:///etc/passwd", "ftp://chai.example/"],
)
async def test_ssrf_and_scheme_guard(fake_dns: dict[str, set[str]], url: str) -> None:
    with pytest.raises(CrawlError):
        await assert_public_url(url)


async def test_crawl_same_domain_robots_and_redirect_ssrf(fake_dns: dict[str, set[str]]) -> None:
    pages = {
        "/robots.txt": (200, "text/plain", "User-agent: *\nDisallow: /private\n"),
        "/": (
            200,
            "text/html",
            '<main><h1>Home</h1><p>Chai and books.</p></main><a href="/about">a</a><a href="/private/x">p</a><a href="https://other.example/">o</a><a href="/sneaky">s</a><a href="/menu.pdf">pdf</a>',
        ),
        "/about": (200, "text/html", "<main><h1>About</h1><p>Founded in 2019.</p></main>"),
        "/sneaky": (302, "text/html", ""),
    }
    fetched: list[str] = []

    def handler(req: httpx.Request) -> httpx.Response:
        fetched.append(f"{req.url.host}{req.url.path}")
        if req.url.path == "/sneaky":
            return httpx.Response(302, headers={"location": "http://evil.example/internal"})
        status, ctype, body = pages.get(req.url.path, (404, "text/html", "nope"))
        return httpx.Response(status, headers={"content-type": ctype}, text=body)

    result = await crawl("https://chai.example/", delay_s=0, transport=httpx.MockTransport(handler))
    assert result.pages == ["https://chai.example/", "https://chai.example/about"]
    assert "disallowed by robots.txt" in result.skipped["https://chai.example/private/x"]
    assert "private or local network" in result.skipped["https://chai.example/sneaky"]
    assert not any(
        "other.example" in f or "evil.example" in f or f.endswith(".pdf") for f in fetched
    )
    assert {s.heading for s in result.sections} == {"Home", "About"}


async def test_crawl_respects_max_pages(fake_dns: dict[str, set[str]]) -> None:
    def handler(req: httpx.Request) -> httpx.Response:
        if req.url.path == "/robots.txt":
            return httpx.Response(404)
        n = int(req.url.path.strip("/") or 0)
        body = f'<main><p>Page {n} about chai.</p></main><a href="/{n + 1}">next</a><a href="/{n + 2}">n2</a>'
        return httpx.Response(200, headers={"content-type": "text/html"}, text=body)

    result = await crawl(
        "https://chai.example/", max_pages=3, delay_s=0, transport=httpx.MockTransport(handler)
    )
    assert len(result.pages) == 3
