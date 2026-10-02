"""Turn brand documents into sections of text with provenance (page, heading)."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from bs4 import BeautifulSoup, Tag


@dataclass
class Section:
    text: str
    heading: str | None = None
    page: int | None = None
    source: str = ""


class ExtractionError(Exception):
    pass


_MD_HEADING = re.compile(r"^(#{1,4})\s+(.+?)\s*#*\s*$")


def from_text(text: str, source: str, *, markdown: bool) -> list[Section]:
    if not markdown:
        return [Section(text=text.strip(), source=source)] if text.strip() else []
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []

    def flush() -> None:
        body = "\n".join(buf).strip()
        if body:
            sections.append(Section(text=body, heading=heading, source=source))
        buf.clear()

    for line in text.splitlines():
        m = _MD_HEADING.match(line)
        if m:
            flush()
            heading = m.group(2).strip()
        else:
            buf.append(line)
    flush()
    return sections


def from_pdf(data: bytes, source: str) -> list[Section]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ExtractionError("This PDF is password-protected. Upload an unlocked copy.")
        pages = [
            (i + 1, (page.extract_text() or "").strip()) for i, page in enumerate(reader.pages)
        ]
    except PdfReadError as exc:
        raise ExtractionError(f"Couldn't read this PDF: {exc}") from exc
    sections = [Section(text=t, page=n, source=source) for n, t in pages if t]
    if not sections:
        raise ExtractionError(
            "No selectable text found. Scanned PDFs need OCR first; export a text PDF instead."
        )
    return sections


def from_docx(data: bytes, source: str) -> list[Section]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise ExtractionError("Couldn't read this Word document.") from exc
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []
    for para in document.paragraphs:
        style = (para.style.name if para.style is not None else "") or ""
        text = para.text.strip()
        if not text:
            continue
        if style.lower().startswith(("heading", "title")):
            if buf:
                sections.append(Section("\n".join(buf), heading, source=source))
                buf = []
            heading = text
        else:
            buf.append(text)
    if buf:
        sections.append(Section("\n".join(buf), heading, source=source))
    for table in document.tables:
        rows = [" | ".join(c.text.strip() for c in row.cells) for row in table.rows]
        if any(r.strip(" |") for r in rows):
            sections.append(Section("\n".join(rows), "Table", source=source))
    return sections


_BOILERPLATE_TAGS = [
    "script",
    "style",
    "noscript",
    "nav",
    "header",
    "footer",
    "aside",
    "form",
    "svg",
    "iframe",
    "button",
]
_HINT_WORDS = (
    "nav|menu|footer|header|cookie|consent|breadcrumb|sidebar"
    "|newsletter|social|share|banner|popup|modal"
)
_BOILERPLATE_HINT = re.compile(rf"(^|[-_ ])({_HINT_WORDS})([-_ ]|$)", re.I)
_BOILERPLATE_ROLES = {"navigation", "banner", "contentinfo", "complementary", "search"}


def from_html(html: str, source: str) -> tuple[list[Section], list[str], str | None]:
    """Return (sections, links, title). Strips nav/footer/cookie boilerplate."""
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else None
    links: list[str] = []
    for a in soup.find_all("a", href=True):
        href = a.get("href") if isinstance(a, Tag) else None
        if isinstance(href, str):
            links.append(href)

    for tag in soup.find_all(_BOILERPLATE_TAGS):
        tag.decompose()
    for tag in soup.find_all(True):
        if not isinstance(tag, Tag) or tag.attrs is None:
            continue
        role = str(tag.get("role") or "").lower()
        ident = " ".join([str(tag.get("id") or ""), *[str(c) for c in (tag.get("class") or [])]])
        if role in _BOILERPLATE_ROLES or (ident and _BOILERPLATE_HINT.search(ident)):
            tag.decompose()

    root = soup.find("main") or soup.find("article") or soup.body or soup
    sections: list[Section] = []
    heading: str | None = None
    buf: list[str] = []
    for el in (
        root.find_all(["h1", "h2", "h3", "p", "li", "blockquote", "td"])
        if isinstance(root, Tag)
        else []
    ):
        text = el.get_text(" ", strip=True)
        if not text:
            continue
        if el.name in ("h1", "h2", "h3"):
            if buf:
                sections.append(Section("\n".join(buf), heading, source=source))
                buf = []
            heading = text
        elif len(text) > 1:
            buf.append(text)
    if buf:
        sections.append(Section("\n".join(buf), heading, source=source))
    return sections, links, title
